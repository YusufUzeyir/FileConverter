import tkinter as tk
from tkinter import filedialog
from pathlib import Path
from fpdf import FPDF
from PIL import Image
from docx import Document
from pdf2docx import Converter
from docx2pdf import convert
import os


def get_download_path():
    """Kullanıcının indirilenler klasörünün yolunu döndürür."""
    home = str(Path.home())
    return os.path.join(home, "Downloads")


def txt_to_pdf(file_path, output_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    with open(file_path, 'r') as file:
        for line in file:
            pdf.cell(0, 10, line, ln=True)
    pdf.output(output_path)


def image_convert(file_path, output_format):
    with Image.open(file_path) as img:
        # PNG dosyasını RGB moduna dönüştür
        rgb_img = img.convert("RGB")
        output_path = f"{get_download_path()}/{Path(file_path).stem}.{output_format}"
        rgb_img.save(output_path)
        return output_path


def docx_to_pdf(file_path, output_path):
    convert(file_path, output_path)


def pdf_to_docx(file_path, output_path):
    cv = Converter(file_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()


def txt_to_docx(file_path, output_path):
    doc = Document()
    with open(file_path, 'r') as file:
        doc.add_paragraph(file.read())
    doc.save(output_path)


def png_to_pdf_or_jpeg_to_pdf(file_path, output_path):
    Image.open(file_path).convert("RGB").save(output_path, "PDF")


def handle_conversion(conversion_type, file_path):
    if conversion_type == "DOCX to PDF":
        output_path = f"{get_download_path()}/{Path(file_path).stem}.pdf"
        docx_to_pdf(file_path, output_path)
    elif conversion_type == "PDF to DOCX":
        output_path = f"{get_download_path()}/{Path(file_path).stem}.docx"
        pdf_to_docx(file_path, output_path)
    elif conversion_type == "JPEG to PNG":
        output_path = image_convert(file_path, "png")
    elif conversion_type == "PNG to JPEG":
        output_path = image_convert(file_path, "jpeg")
    elif conversion_type == "TXT to PDF":
        output_path = f"{get_download_path()}/{Path(file_path).stem}.pdf"
        txt_to_pdf(file_path, output_path)
    elif conversion_type == "TXT to DOCX":
        output_path = f"{get_download_path()}/{Path(file_path).stem}.docx"
        txt_to_docx(file_path, output_path)
    elif conversion_type == "PNG to PDF":
        output_path = f"{get_download_path()}/{Path(file_path).stem}.pdf"
        png_to_pdf_or_jpeg_to_pdf(file_path, output_path)
    elif conversion_type == "JPEG to PDF":
        output_path = f"{get_download_path()}/{Path(file_path).stem}.pdf"
        png_to_pdf_or_jpeg_to_pdf(file_path, output_path)
    else:
        output_path = None
    return output_path


def main():
    root = tk.Tk()
    root.withdraw()  # Tkinter penceresini gizle

    file_path = filedialog.askopenfilename()
    if not file_path:
        print("Dosya seçilmedi.")
        return

    main_window = tk.Tk()
    main_window.title("Dönüşüm Seçin")

    conversion_options = [
        "DOCX to PDF",
        "PDF to DOCX",
        "JPEG to PNG",
        "PNG to JPEG",
        "TXT to PDF",
        "TXT to DOCX",
        "PNG to PDF",
        "JPEG to PDF"
    ]

    top_frame = tk.Frame(main_window)
    bottom_frame = tk.Frame(main_window)
    top_frame.pack(side="top")
    bottom_frame.pack(side="bottom")

    for i, option in enumerate(conversion_options):
        if i < 4:
            frame = top_frame
        else:
            frame = bottom_frame
        button = tk.Button(frame, text=option, command=lambda opt=option: handle_and_print_result(opt, file_path))
        button.pack(side="left", padx=5, pady=5)

    def handle_and_print_result(option, file_path):
        output_path = handle_conversion(option, file_path)
        if output_path:
            print(f"Dosya {output_path} konumuna kaydedildi.")
        else:
            print("Geçersiz seçim.")

    main_window.mainloop()


if __name__ == "__main__":
    main()
